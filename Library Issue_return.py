import tkinter as tk
from tkcalendar import DateEntry
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from tkcalendar import DateEntry
import os


# EXCEL SETTINGS

FILE_NAME = "library_records.xlsx"

COLUMNS = [
    "Book ID",
    "Book Name",
    "Book Author",
    "Student/Person Name",
    "Issue Date",
    "Return Date",
    "Status"
]


# CREATE EXCEL FILE

def create_excel_file():

    if not os.path.exists(FILE_NAME):

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Library Records"

        sheet.append(COLUMNS)

        workbook.save(FILE_NAME)


# TO CLEAR THE FORM

def clear_form():

    book_id_entry.delete(0, tk.END)
    book_name_entry.delete(0, tk.END)
    author_entry.delete(0, tk.END)
    person_entry.delete(0, tk.END)

    action_var.set("Issue")

    issue_date_label.grid_remove()
    issue_date_entry.grid_remove()

    return_date_label.grid_remove()
    return_date_entry.grid_remove()

    save_button.config(text="Save Record")


# FOR CLEAR SCREEN

def clear_screen():

    clear_form()

    form_frame.pack_forget()
    records_frame.pack_forget()

    title_label.config(text="Library Dashboard")


# SHOW DATE FIELD

def show_date_field(event=None):

    selected_action = action_var.get()

    issue_date_label.grid_remove()
    issue_date_entry.grid_remove()

    return_date_label.grid_remove()
    return_date_entry.grid_remove()

    if selected_action == "Issue":

        issue_date_label.grid(row=4, column=0, padx=10, pady=8)

        issue_date_entry.grid(row=4, column=1, padx=10, pady=8)

    elif selected_action == "Return":

        return_date_label.grid(row=4, column=0, padx=10, pady=8)

        return_date_entry.grid(row=4, column=1, padx=10, pady=8)


# ADD RECORD

def show_add_form():

    clear_screen()

    title_label.config(text="Add Book Record")

    form_frame.pack(padx=30, pady=20, fill="x")

    action_label.config(text="Issue / Return")

    action_combo["values"] = [
        "Issue",
        "Return"
    ]

    action_var.set("Issue")

    save_button.config(text="Save Record")

    show_date_field()


# UPDATE / RETURN RECORD

def show_update_form():

    clear_screen()

    title_label.config(text="Update / Return Book")

    form_frame.pack(padx=30, pady=20, fill="x")

    action_label.config(text="Book Action")

    action_combo["values"] = [
        "Return"
    ]

    action_var.set("Return")

    save_button.config(text="Return Book")

    show_date_field()

    records_frame.pack(padx=25, pady=7, fill="both", expand=True)

    load_records()


# DELETE RECORD PAGE

def show_delete_records():

    clear_screen()

    title_label.config(text="Delete Book Record")

    records_frame.pack(padx=30, pady=20, fill="both", expand=True)

    load_records()


# SHOW ALL RECORDS


def show_all_records():

    clear_screen()

    title_label.config(text="All Record Details")

    records_frame.pack(padx=30, pady=20, fill="both", expand=True)

    load_records()


# ADD / RETURN ACTION

def perform_action():

    book_id = book_id_entry.get().strip()
    book_name = book_name_entry.get().strip()
    author = author_entry.get().strip()
    person = person_entry.get().strip()

    selected_action = action_var.get()

    issue_date = issue_date_entry.get().strip()
    return_date = return_date_entry.get().strip()


    # ISSUE BOOK


    if selected_action == "Issue":

        if (
            book_id == ""
            or book_name == ""
            or author == ""
            or person == ""
            or issue_date == ""
        ):

            messagebox.showwarning("Missing Data", "Please fill all required fields.")

            return

        try:

            workbook = load_workbook(FILE_NAME)
            sheet = workbook.active

            # Check duplicate Book ID
            for row in sheet.iter_rows(min_row=2, values_only=True):

                if str(row[0]).lower() == book_id.lower():

                    messagebox.showerror("Duplicate Book ID", "This Book ID already exists.")

                    workbook.close()

                    return


            sheet.append([
                book_id,
                book_name,
                author,
                person,
                issue_date,
                "",
                "Issued"
            ])

            workbook.save(FILE_NAME)
            workbook.close()

            messagebox.showinfo("Success", "Book issued successfully.")

            clear_form()

            show_all_records()

        except Exception as error:

            messagebox.showerror(
                "Error",
                str(error)
            )


    # RETURN BOOK


    elif selected_action == "Return":

        if (
            book_id == ""
            or book_name == ""
            or author == ""
            or person == ""
            or return_date == ""
        ):

            messagebox.showwarning("Missing Data", "Please fill all required fields.")

            return

        try:

            workbook = load_workbook(FILE_NAME)
            sheet = workbook.active

            found = False

            for row in sheet.iter_rows(min_row=2):

                if str(
                    row[0].value
                ).lower() == book_id.lower():

                    row[1].value = book_name
                    row[2].value = author
                    row[3].value = person
                    row[5].value = return_date
                    row[6].value = "Returned"

                    found = True

                    break


            if found:

                workbook.save(FILE_NAME)
                workbook.close()

                messagebox.showinfo("Success", "Book returned successfully.")

                clear_form()

                show_all_records()

            else:

                workbook.close()

                messagebox.showerror("Not Found", "Book ID not found. Enter correct details.")

        except Exception as error:

            messagebox.showerror(
                "Error",
                str(error)
            )


# LOAD RECORDS


def load_records():

    for item in tree.get_children():

        tree.delete(item)

    try:

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

        for row in sheet.iter_rows(min_row=2, values_only=True):

            tree.insert("", tk.END, values=row)

        workbook.close()

    except Exception as error:

        messagebox.showerror(
            "Error",
            str(error)
        )


# SELECT RECORD


def select_record(event=None):

    selected_item = tree.focus()

    if selected_item == "":
        return

    values = tree.item(selected_item, "values")

    if not values:
        return


      
    # DELETE PAGE
      

    if title_label.cget("text") == "Delete Book Record":

        delete_record()

        return


      
    # UPDATE / RETURN PAGE
      

    if title_label.cget("text") == "Update / Return Book":

        clear_form()

        book_id_entry.insert(
            0,
            values[0]
        )

        book_name_entry.insert(
            0,
            values[1]
        )

        author_entry.insert(
            0,
            values[2]
        )

        person_entry.insert(
            0,
            values[3]
        )

        action_var.set("Return")

        show_date_field()

        return


# DELETE RECORD


def delete_record():

    selected_item = tree.focus()

    if selected_item == "":
        return

    values = tree.item(selected_item, "values")

    book_id = values[0]


    answer = messagebox.askyesno(
        "Confirm Delete",
        "Are you sure you want to delete Book ID "
        + str(book_id)
        + "?"
    )


    if answer:

        try:

            workbook = load_workbook(FILE_NAME)
            sheet = workbook.active

            for row in range(2, sheet.max_row + 1):

                if str(
                    sheet.cell(row, 1).value
                ) == str(book_id):

                    sheet.delete_rows(row, 1)

                    break


            workbook.save(FILE_NAME)
            workbook.close()

            messagebox.showinfo("Success", "Record deleted successfully.")

            load_records()

        except Exception as error:

            messagebox.showerror(
                "Error",
                str(error)
            )


# LOGOUT


def logout():

    answer = messagebox.askyesno("Logout", "Are you sure you want to logout?")

    if answer:

        dashboard_frame.pack_forget()

        login_frame.pack(fill="both", expand=True)

        username_entry.delete(0, tk.END)

        password_entry.delete(0, tk.END)


# LOGIN


def login():

    username = username_entry.get().strip()

    password = password_entry.get().strip()


    if (
        username == "admin"
        and password == "admin123"
    ):

        login_frame.pack_forget()

        dashboard_frame.pack(fill="both", expand=True)

    else:

        messagebox.showerror("Login Failed", "Invalid username or password.")


# MAIN WINDOW


root = tk.Tk()

root.title("Library Book Issue & Return System")

root.geometry("1100x700")

root.resizable(False, False)

root.configure(bg="#F8F6F2")


# LOGIN PAGE


login_frame = tk.Frame(root, bg="#F8F6F2")

login_frame.pack(fill="both", expand=True)


login_title = tk.Label(
    login_frame,
    text="LIBRARY BOOK ISSUE & RETURN SYSTEM",
    font=("Arial", 20, "bold"),
    bg="#F8F6F2",
    fg="#555577"
)

login_title.pack(pady=60)


login_box = tk.Frame(login_frame, bg="#EEEAF7", padx=35, pady=30)

login_box.pack()


username_label = tk.Label(
    login_box,
    text="Username",
    bg="#EEEAF7",
    font=("Arial", 11)
)

username_label.grid(row=0, column=0, padx=10, pady=10)


username_entry = tk.Entry(login_box, width=30)

username_entry.grid(row=0, column=1, padx=10, pady=10)


password_label = tk.Label(
    login_box,
    text="Password",
    bg="#EEEAF7",
    font=("Arial", 11)
)

password_label.grid(row=1, column=0, padx=10, pady=10)


password_entry = tk.Entry(login_box, width=30, show="*")

password_entry.grid(row=1, column=1, padx=10, pady=10)


login_button = tk.Button(login_box, text="Login", width=15, bg="#B9DDF2", activebackground="#A8D2EA", command=login)

login_button.grid(row=2, column=0, columnspan=2, pady=20)


login_info = tk.Label(login_frame, text="Username: admin ..... Password: admin123", bg="#F8F6F2", fg="#777777")

login_info.pack()


# DASHBOARD


dashboard_frame = tk.Frame(root, bg="#F8F6F2")


# HEADING


heading_frame = tk.Frame(dashboard_frame, bg="#F8F6F2")

heading_frame.pack(fill="x", padx=30, pady=20)


dashboard_heading = tk.Label(
    heading_frame,
    text="LIBRARY BOOK ISSUE & RETURN SYSTEM",
    font=("Arial", 20, "bold"),
    bg="#F8F6F2",
    fg="#555577"
)

dashboard_heading.pack(side="left")


logout_button = tk.Button(heading_frame, text="Logout", width=15, bg="#B9DDF2", activebackground="#A8D2EA", command=logout)

logout_button.pack(side="right")


# TITLE


title_label = tk.Label(
    dashboard_frame,
    text="Library Dashboard",
    font=("Arial", 15, "bold"),
    bg="#F8F6F2",
    fg="#555577"
)

title_label.pack(pady=5)


# MAIN MENU


menu_frame = tk.Frame(dashboard_frame, bg="#F8F6F2")

menu_frame.pack(pady=20)


add_menu_button = tk.Button(
    menu_frame,
    text="Add Record / Issue Book",
    width=22,
    height=2,
    bg="#DDEFD8",
    activebackground="#CBE6C4",
    command=show_add_form
)

add_menu_button.grid(row=0, column=0, padx=8, pady=8)


update_menu_button = tk.Button(
    menu_frame,
    text="Update / Return Book",
    width=22,
    height=2,
    bg="#E4DDF5",
    activebackground="#D5CBEF",
    command=show_update_form
)

update_menu_button.grid(row=0, column=1, padx=8, pady=8)


delete_menu_button = tk.Button(
    menu_frame,
    text="Delete Record",
    width=18,
    height=2,
    bg="#F6D6D6",
    activebackground="#EBC3C3",
    command=show_delete_records
)

delete_menu_button.grid(row=0, column=2, padx=8, pady=8)


clear_menu_button = tk.Button(
    menu_frame,
    text="Clear / Reset",
    width=18,
    height=2,
    bg="#FFF0C7",
    activebackground="#F6E4B3",
    command=clear_screen
)

clear_menu_button.grid(row=0, column=3, padx=8, pady=8)


show_menu_button = tk.Button(
    menu_frame,
    text="Show All Records",
    width=18,
    height=2,
    bg="#D9EEF7",
    activebackground="#C5E3F0",
    command=show_all_records
)

show_menu_button.grid(row=0, column=4, padx=8, pady=8)


# FORM


form_frame = tk.LabelFrame(
    dashboard_frame,
    text="Book Issue / Return Details",
    font=("Arial", 14, "bold"),
    bg="#FFF8E8",
    fg="#555577",
    padx=20,
    pady=15
)


# BOOK ID


book_id_label = tk.Label(form_frame, text="Book ID", bg="#FFF8E8")

book_id_label.grid(row=0, column=0, padx=10, pady=8)


book_id_entry = tk.Entry(form_frame, width=30)

book_id_entry.grid(row=0, column=1, padx=10, pady=8)


# BOOK NAME


book_name_label = tk.Label(form_frame, text="Book Name", bg="#FFF8E8")

book_name_label.grid(row=1, column=0, padx=10, pady=8)


book_name_entry = tk.Entry(form_frame, width=30)

book_name_entry.grid(row=1, column=1, padx=10, pady=8)


# BOOK AUTHOR


author_label = tk.Label(form_frame, text="Book Author", bg="#FFF8E8")

author_label.grid(row=2, column=0, padx=10, pady=8)


author_entry = tk.Entry(form_frame, width=30)

author_entry.grid(row=2, column=1, padx=10, pady=8)


# STUDENT / PERSON


person_label = tk.Label(form_frame, text="Student / Person Name", bg="#FFF8E8")

person_label.grid(row=3, column=0, padx=10, pady=8)


person_entry = tk.Entry(form_frame, width=30)

person_entry.grid(row=3, column=1, padx=10, pady=8)


# ACTION


action_label = tk.Label(form_frame, text="Issue / Return", bg="#FFF8E8")

action_label.grid(row=0, column=2, padx=10, pady=8)


action_var = tk.StringVar()

action_var.set("Issue")


action_combo = ttk.Combobox(
    form_frame,
    textvariable=action_var,
    values=[
        "Issue",
        "Return"
    ],
    state="readonly",
    width=27
)

action_combo.grid(row=0, column=3, padx=10, pady=8)


action_combo.bind("<<ComboboxSelected>>", show_date_field)


# ISSUE DATE


issue_date_label = tk.Label(form_frame, text="Issue Date", bg="#FFF8E8")


issue_date_entry = DateEntry(
    form_frame,
    width=27,
    date_pattern="dd-mm-yyyy",
    background="#B9DDF2",
    foreground="black",
    borderwidth=2
)


# RETURN DATE


return_date_label = tk.Label(form_frame, text="Return Date", bg="#FFF8E8")


return_date_entry = DateEntry(
    form_frame,
    width=27,
    date_pattern="dd-mm-yyyy",
    background="#B9DDF2",
    foreground="black",
    borderwidth=2
)


# SAVE / RETURN BUTTON


save_button = tk.Button(
    form_frame,
    text="Save Record",
    width=18,
    bg="#DDEFD8",
    activebackground="#CBE6C4",
    command=perform_action
)

save_button.grid(row=5, column=0, columnspan=4, pady=15)


# RECORDS FRAME


records_frame = tk.LabelFrame(
    dashboard_frame,
    text="All Record Details",
    font=("Arial", 14, "bold"),
    bg="#F1F7FB",
    fg="#555577",
    padx=10,
    pady=10
)


columns = (
    "Book ID",
    "Book Name",
    "Book Author",
    "Student/Person",
    "Issue Date",
    "Return Date",
    "Status"
)


# TREEVIEW STYLE


style = ttk.Style()

style.configure("Treeview.Heading",font=("Arial", 11, "bold"))

style.configure("Treeview",font=("Arial", 10),rowheight=28)
tree = ttk.Treeview(records_frame, columns=columns, show="headings", height=20)

for column in columns:

    tree.heading(column, text=column)

    tree.column(column, width=135, anchor="center")


tree.pack(side="left", fill="both", expand=True)


scrollbar = ttk.Scrollbar(records_frame, orient="vertical", command=tree.yview)

scrollbar.pack(side="right", fill="y")


tree.configure(yscrollcommand=scrollbar.set)


tree.bind("<ButtonRelease-1>", select_record)


# START PROGRAM From HERE


create_excel_file()

root.mainloop()
